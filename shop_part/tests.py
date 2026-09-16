import csv
from decimal import Decimal
from io import BytesIO, StringIO

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from tablib import Dataset
from import_export.formats.base_formats import XLSX

from .admin import ProductResource
from .models import Category, Format, Manufactor, Product


class ProductPriceTests(TestCase):
    def test_legacy_product_without_price_remains_valid(self):
        category = Category.objects.create(name="Кирпич")
        product = Product(name="Товар без цены", category=category)
        product.full_clean()
        product.save()

        response = self.client.get(reverse("api-product-detail", args=[product.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertIsNone(response.json()["price"])
        self.assertIsNone(response.json()["old_price"])
        self.assertEqual(response.json()["currency"], "RUB")

    def test_old_price_must_be_higher_than_current_price(self):
        product = Product(
            name="Некорректная скидка",
            price=Decimal("100.00"),
            old_price=Decimal("90.00"),
        )

        with self.assertRaises(ValidationError) as error:
            product.full_clean()

        self.assertIn("old_price", error.exception.message_dict)

    def test_old_price_requires_current_price(self):
        product = Product(name="Нет актуальной цены", old_price=Decimal("150.00"))

        with self.assertRaises(ValidationError) as error:
            product.full_clean()

        self.assertIn("old_price", error.exception.message_dict)


class YandexFeedTests(TestCase):
    expected_header = [
        "ID", "ID2", "URL", "Image", "Title", "Description",
        "Price", "Currency", "Old Price",
        "custom_label_0", "custom_label_1", "custom_label_2",
        "custom_label_3", "custom_label_4", "custom_score",
    ]

    def parse_feed(self):
        response = self.client.get(reverse("yandex_feed"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        return list(csv.reader(StringIO(response.content.decode("utf-8"))))

    def test_feed_contains_utf8_prices_and_escaped_text(self):
        product = Product.objects.create(
            name='Кирпич "Красный", гладкий',
            description="<p>Описание, товара</p>\nвторая строка",
            card_image="products/red-brick.jpg",
            price=Decimal("1250.50"),
            old_price=Decimal("1400.00"),
            currency="RUB",
            feed_id2="RED-01",
            custom_label_0="Акция",
            custom_score=7,
        )

        rows = self.parse_feed()

        self.assertEqual(rows[0], self.expected_header)
        self.assertEqual(len(rows), 2)
        row = dict(zip(rows[0], rows[1]))
        self.assertEqual(row["ID"], str(product.pk))
        self.assertEqual(row["ID2"], "RED-01")
        self.assertEqual(row["Title"], 'Кирпич "Красный", гладкий')
        self.assertEqual(row["Description"], "Описание, товара вторая строка")
        self.assertEqual(row["Price"], "1250.50")
        self.assertEqual(row["Currency"], "RUB")
        self.assertEqual(row["Old Price"], "1400.00")
        self.assertEqual(row["custom_score"], "7")
        self.assertTrue(row["URL"].endswith(f"/p/{product.pk}/"))
        self.assertTrue(row["Image"].endswith("/uploads/products/red-brick.jpg"))

    def test_feed_keeps_price_columns_empty_for_legacy_product(self):
        Product.objects.create(name="Старый товар без цены")

        rows = self.parse_feed()
        row = dict(zip(rows[0], rows[1]))

        self.assertEqual(row["Price"], "")
        self.assertEqual(row["Currency"], "")
        self.assertEqual(row["Old Price"], "")

    def test_disabled_product_is_not_exported(self):
        Product.objects.create(name="Не выгружать", feed_enabled=False)

        rows = self.parse_feed()

        self.assertEqual(rows, [self.expected_header])


class ProductDetailPageTests(TestCase):
    def test_price_block_and_calculation_button_replace_old_callback_block(self):
        product = Product.objects.create(name="Товар")

        response = self.client.get(reverse("product_client_view", args=[product.pk]))
        html = response.content.decode("utf-8")

        self.assertContains(response, "Запросить расчет")
        self.assertContains(response, 'id="price"')
        self.assertNotIn("Узнать цену", html)
        self.assertNotIn("Заказать обратный звонок", html)
        self.assertNotIn("Заказать звонок", html)


class ProductExcelImportTests(TestCase):
    def setUp(self):
        self.category = Category.objects.create(name="Кирпич")
        self.manufacturer = Manufactor.objects.create(name="ЛСР")
        self.product_format = Format.objects.create(name="250×120×65")
        self.product = Product.objects.create(
            name="Исходное название",
            category=self.category,
            manufacturer=self.manufacturer,
            card_image="products/existing.jpg",
            price=Decimal("100.00"),
        )

    def dataset_from_rows(self, *rows):
        resource = ProductResource()
        headers = [field.column_name for field in resource.get_export_fields()]
        dataset = Dataset(headers=headers)
        for row in rows:
            dataset.append([row.get(header, "") for header in headers])
        return dataset

    def exported_product_row(self):
        resource = ProductResource()
        return resource.export(Product.objects.filter(pk=self.product.pk)).dict[0]

    def test_dry_run_then_atomic_update_preserves_photo_and_clears_blank_field(self):
        row = self.exported_product_row()
        row.update({
            "Название товара": "Обновлённое название",
            "Производитель": "",
            "Форматы": self.product_format.name,
            "Актуальная цена": "125.50",
        })
        dataset = self.dataset_from_rows(row)
        resource = ProductResource()

        preview = resource.import_data(
            dataset,
            dry_run=True,
            use_transactions=True,
            rollback_on_validation_errors=True,
        )
        self.assertFalse(preview.has_errors())
        self.assertFalse(preview.has_validation_errors())
        self.product.refresh_from_db()
        self.assertEqual(self.product.name, "Исходное название")

        result = resource.import_data(
            dataset,
            dry_run=False,
            use_transactions=True,
            rollback_on_validation_errors=True,
        )
        self.assertFalse(result.has_errors())
        self.assertFalse(result.has_validation_errors())

        self.product.refresh_from_db()
        self.assertEqual(self.product.name, "Обновлённое название")
        self.assertIsNone(self.product.manufacturer)
        self.assertEqual(self.product.price, Decimal("125.50"))
        self.assertEqual(self.product.card_image.name, "products/existing.jpg")
        self.assertEqual(list(self.product.formats.values_list("name", flat=True)), ["250×120×65"])

    def test_blank_id_creates_new_product(self):
        row = {
            "ID": "",
            "Название товара": "Новый товар",
            "Категория": self.category.name,
            "Валюта": "RUB",
            "Приоритет": 0,
            "Включать в Яндекс-фид": "Да",
        }
        result = ProductResource().import_data(
            self.dataset_from_rows(row),
            dry_run=False,
            use_transactions=True,
            rollback_on_validation_errors=True,
        )

        self.assertFalse(result.has_errors())
        self.assertFalse(result.has_validation_errors())
        self.assertTrue(Product.objects.filter(name="Новый товар").exists())

    def test_invalid_row_rolls_back_valid_row(self):
        valid_row = self.exported_product_row()
        valid_row["Название товара"] = "Не должно сохраниться"
        invalid_row = dict(valid_row)
        invalid_row.update({
            "ID": "",
            "Название товара": "Товар с ошибкой",
            "Категория": "Несуществующая категория",
        })

        result = ProductResource().import_data(
            self.dataset_from_rows(valid_row, invalid_row),
            dry_run=False,
            use_transactions=True,
            rollback_on_validation_errors=True,
        )

        self.assertTrue(result.has_errors() or result.has_validation_errors())
        self.product.refresh_from_db()
        self.assertEqual(self.product.name, "Исходное название")
        self.assertFalse(Product.objects.filter(name="Товар с ошибкой").exists())

    def test_unknown_existing_id_is_rejected(self):
        row = self.exported_product_row()
        row["ID"] = 999999

        result = ProductResource().import_data(
            self.dataset_from_rows(row),
            dry_run=True,
            use_transactions=True,
            rollback_on_validation_errors=True,
        )

        self.assertTrue(result.has_errors() or result.has_validation_errors())


class ProductExcelWorkbookTests(TestCase):
    def test_admin_download_contains_products_references_and_dropdowns_without_photos(self):
        category = Category.objects.create(name="Кирпич")
        Product.objects.create(
            name="Товар для выгрузки",
            category=category,
            card_image="products/hidden-from-excel.jpg",
        )
        user = get_user_model().objects.create_superuser(
            username="excel-admin",
            email="admin@example.com",
            password="test-password",
        )
        self.client.force_login(user)

        response = self.client.get(reverse("admin:shop_part_product_download_update_excel"))

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["Товары", "Справочники", "Инструкция"])
        self.assertEqual(workbook.active.title, "Товары")

        products_sheet = workbook["Товары"]
        headers = [cell.value for cell in products_sheet[1]]
        self.assertIn("Название товара", headers)
        self.assertIn("Категория", headers)
        self.assertNotIn("Главное фото", headers)
        self.assertGreaterEqual(len(products_sheet.data_validations.dataValidation), 10)
        all_values = [cell.value for row in products_sheet.iter_rows() for cell in row]
        self.assertNotIn("products/hidden-from-excel.jpg", all_values)

        imported_dataset = XLSX().create_dataset(response.content)
        preview = ProductResource().import_data(
            imported_dataset,
            dry_run=True,
            use_transactions=True,
            rollback_on_validation_errors=True,
        )
        self.assertFalse(preview.has_errors())
        self.assertFalse(preview.has_validation_errors())
        self.assertEqual(preview.totals["skip"], 1)
        self.assertEqual(preview.totals["update"], 0)

    def test_product_admin_changelist_renders_excel_download_button(self):
        user = get_user_model().objects.create_superuser(
            username="product-admin",
            email="product-admin@example.com",
            password="test-password",
        )
        self.client.force_login(user)

        response = self.client.get(reverse("admin:shop_part_product_changelist"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Скачать Excel для обновления")

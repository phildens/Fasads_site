from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter, quote_sheetname


HEADER_FILL = PatternFill("solid", fgColor="422223")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ID_FILL = PatternFill("solid", fgColor="E7E3E2")
INPUT_FILL = PatternFill("solid", fgColor="FFF8E8")
ERROR_FILL = PatternFill("solid", fgColor="FCE8E6")
THIN_BORDER = Border(bottom=Side(style="thin", color="D9D2D0"))


def _add_reference_sheet(workbook, reference_lists):
    sheet = workbook.create_sheet("Справочники")
    sheet.sheet_view.showGridLines = False

    range_names = {}
    for column_index, (header, values) in enumerate(reference_lists.items(), start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

        clean_values = [str(value).strip() for value in values if str(value).strip()]
        for row_index, value in enumerate(clean_values, start=2):
            sheet.cell(row=row_index, column=column_index, value=value)

        last_row = max(2, len(clean_values) + 1)
        range_name = f"reference_{column_index}"
        reference = (
            f"{quote_sheetname(sheet.title)}!"
            f"${get_column_letter(column_index)}$2:${get_column_letter(column_index)}${last_row}"
        )
        workbook.defined_names.add(DefinedName(range_name, attr_text=reference))
        range_names[header] = range_name
        sheet.column_dimensions[get_column_letter(column_index)].width = max(
            16,
            min(42, max([len(header), *[len(value) for value in clean_values]]) + 2),
        )

    sheet.freeze_panes = "A2"
    return range_names


def _add_instructions_sheet(workbook):
    sheet = workbook.create_sheet("Инструкция")
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "Обновление товаров через Excel"
    sheet["A1"].font = Font(size=16, bold=True, color="422223")

    instructions = [
        "1. Редактируйте данные только на листе «Товары».",
        "2. Не изменяйте ID существующего товара. Пустой ID создаёт новый товар.",
        "3. Для справочных полей выбирайте значения из выпадающих списков.",
        "4. Несколько форматов или значений пустотности разделяйте точкой с запятой (;).",
        "5. Пустая ячейка очищает необязательное поле.",
        "6. Старая цена должна быть выше актуальной цены.",
        "7. Фотографии в Excel не загружаются. Существующие фотографии сохраняются.",
        "8. При импорте сначала проверьте предварительный результат, затем подтвердите обновление.",
        "9. Если хотя бы одна строка содержит ошибку, изменения не применяются.",
        "10. Товары, которых нет в файле, остаются на сайте.",
    ]
    for row_index, text in enumerate(instructions, start=3):
        sheet.cell(row=row_index, column=1, value=text)
        sheet.cell(row=row_index, column=1).alignment = Alignment(wrap_text=True, vertical="top")

    sheet.column_dimensions["A"].width = 105
    for row_index in range(3, 3 + len(instructions)):
        sheet.row_dimensions[row_index].height = 28


def build_products_workbook(dataset, reference_lists):
    """Build the employee-facing product workbook with live reference dropdowns."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Товары"
    sheet.sheet_view.showGridLines = False

    headers = list(dataset.headers)
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    rows = list(dataset)
    if not rows:
        rows = [[None] * len(headers)]

    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = THIN_BORDER
            if column_index == 1:
                cell.fill = ID_FILL
            else:
                cell.fill = INPUT_FILL

    header_to_column = {header: index for index, header in enumerate(headers, start=1)}
    for header in ("Актуальная цена", "Старая цена"):
        if header in header_to_column:
            column_letter = get_column_letter(header_to_column[header])
            for cell in sheet[column_letter][1:]:
                cell.number_format = "#,##0.00"

    for header in ("ID", "Приоритет", "custom_score"):
        if header in header_to_column:
            column_letter = get_column_letter(header_to_column[header])
            for cell in sheet[column_letter][1:]:
                cell.number_format = "0"

    widths = {
        "ID": 9,
        "Название товара": 34,
        "Категория": 24,
        "Производитель": 24,
        "Вид материала": 22,
        "Тип товара": 22,
        "Цвет": 20,
        "Форматы": 28,
        "Пустотность": 22,
        "Морозостойкость": 20,
        "Марка прочности": 20,
        "Водопоглощение": 20,
        "Описание": 52,
        "Актуальная цена": 17,
        "Старая цена": 15,
        "Валюта": 12,
        "Метка товара": 17,
        "Приоритет": 12,
        "Включать в Яндекс-фид": 22,
        "ID2": 16,
        "Прайс-лист товара": 24,
    }
    for column_index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = widths.get(header, 20)

    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    sheet.row_dimensions[1].height = 42
    for row_index in range(2, len(rows) + 2):
        sheet.row_dimensions[row_index].height = 22

    table = Table(
        displayName="ProductsTable",
        ref=f"A1:{get_column_letter(len(headers))}{len(rows) + 1}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)

    range_names = _add_reference_sheet(workbook, reference_lists)
    multi_value_headers = {"Форматы", "Пустотность"}
    for header, range_name in range_names.items():
        column_index = header_to_column.get(header)
        if not column_index:
            continue
        validation = DataValidation(
            type="list",
            formula1=f"={range_name}",
            allow_blank=True,
        )
        validation.promptTitle = header
        if header in multi_value_headers:
            validation.prompt = "Можно указать несколько значений через точку с запятой (;)."
            validation.showErrorMessage = False
        else:
            validation.errorTitle = "Значение отсутствует в справочнике"
            validation.error = "Выберите значение из выпадающего списка."
            validation.showErrorMessage = True
        validation.showInputMessage = True
        sheet.add_data_validation(validation)
        column_letter = get_column_letter(column_index)
        validation.add(f"{column_letter}2:{column_letter}5000")

    if "Старая цена" in header_to_column and "Актуальная цена" in header_to_column:
        old_letter = get_column_letter(header_to_column["Старая цена"])
        price_letter = get_column_letter(header_to_column["Актуальная цена"])
        sheet.conditional_formatting.add(
            f"{old_letter}2:{old_letter}5000",
            FormulaRule(
                formula=[f'AND(${old_letter}2<>"",OR(${price_letter}2="",${old_letter}2<=${price_letter}2))'],
                fill=ERROR_FILL,
            ),
        )

    _add_instructions_sheet(workbook)
    workbook.active = 0

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
